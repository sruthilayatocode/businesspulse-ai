from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\SURESH KUMAR\Documents\Codex\2026-09-08\news-to-business-impact-agent-tracks\outputs\BusinessPulse_AI_PRD.docx"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F5F8FC"
GRAY = "D9D9D9"
TEXT = RGBColor(0, 0, 0)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_border(cell, color=GRAY):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def keep_with_next(p):
    pPr = p._p.get_or_add_pPr()
    node = OxmlElement('w:keepNext')
    pPr.append(node)

def add_text(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    r.font.name = 'Aptos'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    r.font.color.rgb = color or TEXT
    return r

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text)
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ''
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, bold=True, size=9, color=RGBColor(255,255,255))
        if widths: cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ''
            set_cell_border(cell)
            set_cell_margins(cell)
            if ridx % 2 == 1: set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(value) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, size=9)
            if widths: cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 12)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = TEXT
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
styles['Heading 1'].paragraph_format.space_before = Pt(16); styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].paragraph_format.space_before = Pt(10); styles['Heading 2'].paragraph_format.space_after = Pt(5)

# cover
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(p, 'BusinessPulse AI', bold=True, size=26)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
add_text(p, 'Product Requirements Document', size=15)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(24)
add_text(p, 'External intelligence connected to internal business exposure', size=11)

add_table(doc, ['Document control', 'Value'], [
    ['Product', 'BusinessPulse AI'],
    ['Purpose', 'Define the MVP for an executive business impact intelligence platform'],
    ['Primary platform', 'Exasol analytics database with AI services and a web dashboard'],
    ['Audience', 'Product, data engineering, AI engineering, business risk, and executive stakeholders'],
    ['Status', 'Draft for project implementation'],
], [1.65, 4.85])

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(15)
add_text(p, 'Executive summary. ', bold=True)
add_text(p, 'BusinessPulse AI continuously monitors external news, regulations, supplier disruptions, market signals, and competitor activity. It links those events to supplier, product, sales, inventory, and market data stored in Exasol, then delivers prioritized, explainable recommendations to decision makers. The MVP focuses on helping a business see what changed, what is exposed, and what to do next.')

doc.add_page_break()

def h1(text):
    p = doc.add_paragraph(style='Heading 1'); add_text(p, text, bold=True, size=15); keep_with_next(p)
def h2(text):
    p = doc.add_paragraph(style='Heading 2'); add_text(p, text, bold=True, size=12); keep_with_next(p)
def para(text):
    p = doc.add_paragraph(); add_text(p, text); return p

h1('1 Product vision and problem')
para('Business leaders receive large volumes of external information but lack a fast, reliable way to connect a news event to their specific operational or financial exposure. Teams often discover disruptions after a supplier misses a delivery, a competitor changes price, or a new regulation is already close to enforcement.')
h2('Product vision')
para('BusinessPulse AI will be the business early-warning layer: a single place where external signals are translated into clear exposure, severity, evidence, and recommended action.')
h2('Problem statements')
add_bullet(doc, 'External news is fragmented and too broad for executives to monitor manually.')
add_bullet(doc, 'Supplier, product, inventory, sales, and market data are often analyzed separately.')
add_bullet(doc, 'Risk alerts rarely quantify business impact or explain the path from event to exposure.')
add_bullet(doc, 'Leadership needs concise, actionable briefings instead of article-level summaries.')

h1('2 Goals and non goals')
add_table(doc, ['Goals for the MVP', 'Non goals for the MVP'], [
    ['Detect and classify events relevant to suppliers, markets, competitors, and regulation.', 'Provide autonomous purchasing, legal, or pricing decisions.'],
    ['Map news entities to an internal company profile and exposure records.', 'Guarantee predictive accuracy for all geopolitical or market events.'],
    ['Prioritize events using transparent risk and business-exposure scoring.', 'Replace the company ERP, supplier-management, or compliance system.'],
    ['Deliver a daily briefing and critical-event alerts with clear next actions.', 'Ingest every global news source or support every industry at launch.'],
], [3.25, 3.25])

h1('3 Users and key use cases')
add_table(doc, ['User', 'Need', 'Primary workflow'], [
    ['Executive', 'Know the few events that require attention.', 'Read daily briefing; open critical alert; approve response.'],
    ['Supply chain manager', 'Prevent material or logistics disruption.', 'Review supplier exposure; validate inventory; assign mitigation.'],
    ['Market and strategy analyst', 'Track demand shifts and competitor moves.', 'Filter by market or competitor; compare event trends.'],
    ['Compliance lead', 'Prepare for regulatory changes.', 'Review affected regulations; track effective dates and owners.'],
    ['Data administrator', 'Maintain trusted inputs and mappings.', 'Load ERP data; manage source reliability and entity mappings.'],
], [1.35, 2.25, 2.9])

h1('4 MVP scope and functional requirements')
h2('4.1 News and event intelligence')
add_table(doc, ['ID', 'Requirement', 'Acceptance criteria'], [
    ['FR 01', 'Ingest selected RSS, news API, and regulatory sources on a scheduled basis.', 'Each item is stored with source, URL, publication time, title, body, and ingest time.'],
    ['FR 02', 'Deduplicate semantically similar news items.', 'Repeated coverage of the same event is grouped into one event cluster.'],
    ['FR 03', 'Classify event type and extract entities.', 'System returns event category, companies, countries, products, commodities, regulations, and confidence.'],
    ['FR 04', 'Support supplier, market, competitor, and regulatory event categories.', 'Dashboard filters and briefing sections use all four categories.'],
], [0.65, 2.65, 3.2])
h2('4.2 Exposure and risk analysis')
add_table(doc, ['ID', 'Requirement', 'Acceptance criteria'], [
    ['FR 05', 'Maintain a company exposure profile in Exasol.', 'Profile includes suppliers, sites, materials, products, markets, competitors, and regulations.'],
    ['FR 06', 'Match extracted entities to internal entities and relationships.', 'Every material event shows linked company records and matching confidence.'],
    ['FR 07', 'Calculate a risk priority score.', 'Score uses severity, likelihood, internal exposure, urgency, and source confidence.'],
    ['FR 08', 'Show the explainable impact chain.', 'User can trace event to entity, exposure, affected product or market, and recommended action.'],
], [0.65, 2.65, 3.2])
h2('4.3 Briefing, alerts, and actions')
add_table(doc, ['ID', 'Requirement', 'Acceptance criteria'], [
    ['FR 09', 'Generate a daily executive briefing.', 'Briefing lists top changes, risks, estimated exposure, and recommended actions.'],
    ['FR 10', 'Create alerts for high and critical scores.', 'Alerts include severity, reason, evidence links, and notification destination.'],
    ['FR 11', 'Allow users to assign and track mitigation actions.', 'Action has owner, due date, status, notes, and related event.'],
    ['FR 12', 'Provide dashboard filters and audit history.', 'Users can filter by region, category, supplier, product, market, and status.'],
], [0.65, 2.65, 3.2])

doc.add_page_break()
h1('5 Risk scoring model')
para('The initial model will be intentionally explainable. Each dimension is normalized from 0 to 100. Business users can change approved weights by industry or category without changing code.')
add_table(doc, ['Dimension', 'Definition', 'Initial weight'], [
    ['Event severity', 'Scale of the reported disruption, policy change, market movement, or competitive move.', '25%'],
    ['Internal exposure', 'Dependency based on supplier criticality, revenue contribution, inventory, or market share.', '35%'],
    ['Urgency', 'Time until business impact or regulatory effective date.', '20%'],
    ['Likelihood', 'Probability the event will materialize or persist, based on evidence and event context.', '10%'],
    ['Source confidence', 'Reliability of the source and AI extraction confidence.', '10%'],
], [1.4, 4.25, 1.0])
para('Priority score = 0.25 x severity + 0.35 x exposure + 0.20 x urgency + 0.10 x likelihood + 0.10 x confidence.')
add_table(doc, ['Score', 'Priority', 'Expected response'], [
    ['80 to 100', 'Critical', 'Immediate alert and named action owner.'],
    ['60 to 79', 'High', 'Review within 24 hours and include in briefing.'],
    ['35 to 59', 'Medium', 'Monitor and include in weekly intelligence view.'],
    ['0 to 34', 'Low', 'Store for history and trend analysis.'],
], [1.2, 1.5, 3.95])

h1('6 Exasol architecture and data design')
para('Exasol is the analytical system of record for internal exposure, normalized events, entity relationships, scoring outputs, and briefing-ready aggregates. The application layer calls Exasol for analytical views and writes approved actions back to Exasol for a complete audit trail.')
h2('Data flow')
add_table(doc, ['Stage', 'Responsibility', 'Exasol role'], [
    ['Collect', 'External ingestion service retrieves RSS, APIs, regulatory feeds, and optional company data.', 'Land raw references and ingestion metadata.'],
    ['Understand', 'AI service extracts event facts, entities, sentiment, and confidence.', 'Persist normalized events and entity mentions.'],
    ['Connect', 'Matching service links external entities to internal exposure records.', 'Join exposure tables and maintain mappings.'],
    ['Prioritize', 'Risk service calculates score and impact estimate.', 'Run SQL analytics and Python UDF scoring where appropriate.'],
    ['Act', 'Dashboard and alert service present results and tasks.', 'Serve briefing, trend, and audit queries.'],
], [1.0, 2.8, 2.85])
h2('Core Exasol tables')
add_table(doc, ['Table', 'Purpose', 'Key fields'], [
    ['RAW_NEWS_ITEMS', 'Raw ingested source records.', 'news_id, source, url, published_at, content_hash'],
    ['NEWS_EVENTS', 'Deduplicated and classified events.', 'event_id, category, severity, event_date, summary, confidence'],
    ['EVENT_ENTITIES', 'Entities extracted from each event.', 'event_id, entity_name, entity_type, canonical_id, confidence'],
    ['COMPANY_EXPOSURE', 'Internal supplier, product, market, and regulatory relationships.', 'entity_id, entity_type, dependency_score, region, revenue_share'],
    ['EVENT_EXPOSURE_MAP', 'Links events to internal exposure.', 'event_id, entity_id, match_type, match_confidence, impact_path'],
    ['RISK_SCORES', 'Scored alerts and financial estimates.', 'event_id, score, priority, revenue_at_risk, calculated_at'],
    ['MITIGATION_ACTIONS', 'User-owned response tracking.', 'action_id, event_id, owner, due_date, status'],
], [1.6, 2.4, 2.65])

h1('7 System requirements')
h2('Non functional requirements')
add_table(doc, ['Area', 'Requirement'], [
    ['Performance', 'Display a filtered risk dashboard within 3 seconds for the MVP data volume.'],
    ['Freshness', 'Ingest configured sources at least every 60 minutes; critical-source polling can be more frequent.'],
    ['Availability', 'Provide scheduled daily briefing even if a non-critical source is temporarily unavailable.'],
    ['Security', 'Use role-based access, encrypted credentials, and least-privilege database roles.'],
    ['Auditability', 'Store event source links, extraction confidence, score inputs, user actions, and score version.'],
    ['Explainability', 'Never show a critical alert without source evidence and an impact explanation.'],
], [1.4, 5.25])

doc.add_page_break()
h1('8 User experience requirements')
h2('Executive home')
add_bullet(doc, 'Top five items by current priority, with a one-sentence impact statement.')
add_bullet(doc, 'A compact risk distribution by supplier, market, competitor, and regulation category.')
add_bullet(doc, 'Daily change summary: new, escalated, resolved, and overdue actions.')
h2('Event detail')
add_bullet(doc, 'Source article and extracted event facts with confidence score.')
add_bullet(doc, 'Impact chain: external event to company exposure to likely business consequence.')
add_bullet(doc, 'Affected suppliers, products, markets, inventory signals, and estimated exposure.')
add_bullet(doc, 'Recommended mitigation actions and current response status.')
h2('Analytics views')
add_bullet(doc, 'Supplier and region risk heat map.')
add_bullet(doc, 'Competitor timeline with product, pricing, partnership, and expansion events.')
add_bullet(doc, 'Regulatory calendar ordered by effective date and business impact.')
add_bullet(doc, 'Trend view for recurring event types, sources, regions, and unresolved risks.')

h1('9 AI design and guardrails')
add_table(doc, ['Capability', 'Implementation approach', 'Guardrail'], [
    ['Extraction', 'LLM converts articles into structured event JSON; rule checks validate required fields.', 'Keep original source and extraction confidence.'],
    ['Entity resolution', 'Hybrid alias dictionary, fuzzy match, and embedding similarity.', 'Human-review queue for ambiguous high-impact matches.'],
    ['Briefing generation', 'LLM summarizes approved high-priority events using Exasol query results.', 'Require evidence links and avoid unsupported financial claims.'],
    ['Recommendations', 'Rule templates plus LLM wording, based on event category and exposure.', 'Mark recommendations as proposed actions, not automatic decisions.'],
], [1.35, 3.1, 2.2])

h1('10 Integrations')
add_table(doc, ['Integration', 'MVP purpose', 'Priority'], [
    ['News APIs and RSS feeds', 'Event source ingestion.', 'Must have'],
    ['Regulatory websites or feeds', 'Compliance and policy monitoring.', 'Must have'],
    ['ERP or procurement export', 'Supplier, material, inventory, and purchase exposure.', 'Must have'],
    ['Sales or CRM export', 'Market and product revenue exposure.', 'Should have'],
    ['Email or Slack or Teams', 'Critical alerts and daily briefing delivery.', 'Should have'],
    ['Identity provider', 'Single sign-on for production usage.', 'Later phase'],
], [1.75, 3.75, 1.15])

h1('11 Delivery plan')
add_table(doc, ['Phase', 'Indicative duration', 'Deliverable'], [
    ['1 Foundation', 'Week 1', 'Exasol schemas, sample company profile, source connector, basic dashboard shell.'],
    ['2 Intelligence', 'Weeks 2 to 3', 'Event extraction, classification, deduplication, entity matching.'],
    ['3 Exposure', 'Weeks 4 to 5', 'Internal data joins, impact chains, explainable risk scoring.'],
    ['4 Action', 'Week 6', 'Briefing generation, alerts, mitigation tracker, executive demo scenario.'],
    ['5 Hardening', 'Week 7', 'User testing, score calibration, security review, performance testing.'],
], [1.25, 1.35, 4.05])

h1('12 Success metrics')
add_table(doc, ['Metric', 'MVP target'], [
    ['Relevant event precision', 'At least 75% of reviewed high-priority alerts are judged relevant by business users.'],
    ['Critical alert latency', 'Critical-source event visible in the dashboard within 60 minutes of ingestion.'],
    ['Explanation completeness', '100% of high and critical alerts show source, exposure mapping, score inputs, and action.'],
    ['Briefing engagement', 'At least 60% of pilot executives open the daily briefing each week.'],
    ['Action closure', 'At least 70% of critical actions have an owner and due date within one business day.'],
], [2.1, 4.55])

h1('13 Risks and mitigation')
add_table(doc, ['Risk', 'Mitigation'], [
    ['False positives create alert fatigue.', 'Use strict priority thresholds, source reliability weighting, and user feedback to tune scores.'],
    ['Entity matching links an event to the wrong supplier or competitor.', 'Use canonical entity lists and require review for ambiguous high-exposure matches.'],
    ['Internal data is incomplete or stale.', 'Display data freshness and confidence; launch with a validated sample profile.'],
    ['LLM output is unsupported or inconsistent.', 'Constrain outputs to structured schemas and show original source evidence.'],
    ['Sensitive supplier and financial data is exposed.', 'Apply role-based access, secure credentials, and data minimization in AI prompts.'],
], [2.55, 4.1])

h1('14 MVP demo scenario')
para('A typhoon disrupts semiconductor production in Taiwan. BusinessPulse AI ingests credible coverage, extracts the disruption and location, and maps it to a company supplier in the affected region. Exasol joins supplier dependency, inventory, bill-of-material, product, and revenue data. The dashboard raises a high-priority alert showing affected products, expected inventory coverage, estimated revenue at risk, and recommended actions: contact the supplier, confirm shipment status, evaluate alternatives, and prioritize high-margin production.')

doc.core_properties.title = 'BusinessPulse AI Product Requirements Document'
doc.core_properties.subject = 'Product requirements for an Exasol based business impact intelligence platform'
doc.core_properties.author = 'BusinessPulse AI Project Team'
doc.save(OUT)
print(OUT)
